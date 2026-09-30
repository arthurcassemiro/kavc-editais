"""Planilha de dados da construção do PCM para o material suplementar (Calibri, paleta dos pilares)."""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.properties import PageSetupProperties
import sys
OUT=sys.argv[1] if len(sys.argv)>1 else "Dados_construcao_PCM.xlsx"
b=pd.read_csv("02_base_tratada/base_municipal.csv").sort_values("mun",key=lambda s:s.str.normalize("NFKD")).reset_index(drop=True)
N=len(b); R0=6; R1=R0+N-1
COR={"FE":("C24C12","FBE3D3"),"FD":("0171C1","DCEAF6"),"FS":("307E3E","E1F1DD"),"PCM":("54278F","ECE8F4"),"BASE":("3F3F3F","EDEDED")}
TXT="262626"; CINZA="7F7F7F"
fino=Side(style="thin",color="D0D0D0"); forte=Side(style="medium",color="595959")
PORTE={"G2":"G2 (acima de 1 milhão)","G1":"G1 (500 mil a 1 milhão)","M2":"M2 (200 a 500 mil)","M1":"M1 (100 a 200 mil)","P4":"P4 (70 a 100 mil)","P3":"P3 (20 a 70 mil)","P2":"P2 (5 a 20 mil)","P1":"P1 (até 5 mil)"}
ORD=["G2","G1","M2","M1","P4","P3","P2","P1"]
wb=Workbook(); wb.remove(wb.active)
def F(size=10,bold=False,color=TXT,italic=False): return Font(name="Calibri",size=size,bold=bold,color=color,italic=italic)
def cab(ws,titulo,sub):
    ws["A1"]=titulo; ws["A1"].font=F(14,True)
    ws["A2"]=sub; ws["A2"].font=F(10,color=CINZA,italic=True)
    ws.sheet_view.showGridLines=False
def grupos(ws,row,specs):
    """specs: lista (col_ini, col_fim, texto, chave_cor)"""
    for c0,c1,t,k in specs:
        if c1>c0: ws.merge_cells(start_row=row,start_column=c0,end_row=row,end_column=c1)
        c=ws.cell(row=row,column=c0,value=t); c.font=F(10,True,"FFFFFF"); c.alignment=Alignment(horizontal="center",vertical="center")
        for cc in range(c0,c1+1):
            ws.cell(row=row,column=cc).fill=PatternFill("solid",fgColor=COR[k][0])
def cabecalho(ws,row,cols,k):
    for j,(t,w) in enumerate(cols,1):
        c=ws.cell(row=row,column=j,value=t); c.font=F(10,True); c.fill=PatternFill("solid",fgColor=COR[k][1])
        c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); c.border=Border(bottom=forte)
        ws.column_dimensions[L(j)].width=w
    ws.row_dimensions[row].height=44
def corpo(ws,c0,c1,formatos,destaque=()):
    for r in range(R0,R1+1):
        for j in range(c0,c1+1):
            c=ws.cell(row=r,column=j); c.font=F(10,bold=(j in destaque)); c.border=Border(bottom=fino)
            if j in formatos: c.number_format=formatos[j]; c.alignment=Alignment(horizontal="right")
def base_cols(ws):
    for i,row in b.iterrows():
        r=R0+i
        ws.cell(row=r,column=1,value=int(row.cod7)); ws.cell(row=r,column=2,value=row.mun); ws.cell(row=r,column=3,value=row.porte_dc)
def rodape(ws,linhas,col=1):
    r=R1+2
    for t in linhas:
        ws.cell(row=r,column=col,value=t).font=F(9,color=CINZA,italic=True); r+=1
def fechar(ws,freeze):
    ws.freeze_panes=freeze
    ws.page_setup.orientation="landscape"; ws.page_setup.paperSize=ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr=PageSetupProperties(fitToPage=True); ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
    ws.print_title_rows=f"{R0-2}:{R0-1}"
# ---------------- Sobre ----------------
ws=wb.create_sheet("Sobre"); ws.sheet_view.showGridLines=False
ws.column_dimensions["A"].width=3; ws.column_dimensions["B"].width=30; ws.column_dimensions["C"].width=92
ws["B2"]="Potencial de Creche por Município (PCM): dados de construção do indicador"; ws["B2"].font=F(15,True)
ws["B3"]="Material suplementar. Dados municipais usados no cálculo do indicador e na distribuição das unidades do Programa Infância Feliz Paraná."; ws["B3"].font=F(10,color=CINZA,italic=True)
linhas=[("Aba","Conteúdo","BASE"),
("1 Indicadores","Nove indicadores brutos dos 399 municípios, com as variáveis de origem, organizados pelos três fatores.","BASE"),
("2 Fator educacional","Indicadores normalizados de 0 a 100 e cálculo do fator educacional (F_E), com fórmulas.","FE"),
("3 Fator demográfico","Indicadores normalizados de 0 a 100 e cálculo do fator demográfico (F_D), com fórmulas.","FD"),
("4 Fator socioeconômico","Indicadores normalizados de 0 a 100 e cálculo do fator socioeconômico (F_S), com fórmulas.","FS"),
("5 PCM","Agregação dos três fatores no PCM, com posição no estado e na faixa de porte.","PCM"),
("6 Distribuição","Creches atribuídas por município em cada versão da distribuição e na lista publicada.","BASE"),
("7 Resumo por porte","Totais por faixa de porte, calculados com fórmulas a partir das abas anteriores.","BASE")]
r=5
for a,t,k in linhas:
    ca=ws.cell(row=r,column=2,value=a); cb=ws.cell(row=r,column=3,value=t)
    if r==5:
        for c in (ca,cb): c.font=F(10,True); c.border=Border(bottom=forte)
    else:
        ca.font=F(10,True,COR[k][0] if k!="BASE" else TXT); cb.font=F(10); ca.border=cb.border=Border(bottom=fino)
    cb.alignment=Alignment(wrap_text=True,vertical="top"); ca.alignment=Alignment(vertical="top")
    r+=1
r+=1
ws.cell(row=r,column=2,value="Método").font=F(11,True); r+=1
met=["Normalização mínimo-máximo de 0 a 100, com extremos dados pelos 399 municípios. Nos indicadores de sentido inverso (cobertura de creche, matrícula na pré-escola e no fundamental, oferta privada e IPDM), o município de menor valor recebe 100.",
"Fator educacional: F_E = normalização de [(2 × MC + MP + OP) / 4]. Fator demográfico: F_D = normalização de [(MI + PC + TN) / 3]. Fator socioeconômico: F_S = normalização de [(BP + CU + IR) / 3].",
"PCM = (2 × F_E + F_D + 2 × F_S) / 500, entre 0 e 1.",
"Os indicadores normalizados das abas 2 a 4 são os valores das planilhas de cálculo do estudo; fatores e PCM são recalculados por fórmula e reproduzem exatamente os valores do estudo.",
"Na planilha de cálculo, a coluna usada como população infantil do indicador PC corresponde, pela magnitude, à faixa de 0 a 9 anos; o ordenamento dos municípios não se altera (correlação de 0,998 com a faixa de 0 a 4 anos)."]
for t in met:
    c=ws.cell(row=r,column=2,value=t); ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=3); c.font=F(10); c.alignment=Alignment(wrap_text=True,vertical="top"); ws.row_dimensions[r].height=30; r+=1
r+=1
ws.cell(row=r,column=2,value="Fontes").font=F(11,True); r+=1
fontes=[("MC, MP, OP","Inep, Censo Escolar 2023; IBGE, Censo Demográfico 2022"),("MI","Ministério da Saúde, SIM e Sinasc (Datasus), 2020"),("PC","IBGE, Censo Demográfico 2022"),("TN","Ipardes, Base de Dados do Estado, 2022"),("BP, CU","Ministério da Saúde, Sisvan e condicionalidades de saúde do Bolsa Família, 2023"),("IR","Ipardes, IPDM 2021, dimensão renda, emprego e produção agropecuária"),("Distribuição","Planilhas de cálculo do estudo (março e maio de 2024) e Resolução SEDEF nº 219/2024")]
for a,t in fontes:
    ws.cell(row=r,column=2,value=a).font=F(10,True); c=ws.cell(row=r,column=3,value=t); c.font=F(10); r+=1
r+=1
ws.cell(row=r,column=2,value="Cores das abas").font=F(11,True); r+=1
for k,t in [("FE","Fator educacional"),("FD","Fator demográfico"),("FS","Fator socioeconômico"),("PCM","Indicador agregado")]:
    c=ws.cell(row=r,column=2,value=t); c.font=F(10,True,"FFFFFF"); c.fill=PatternFill("solid",fgColor=COR[k][0]); r+=1
ws.page_setup.orientation="landscape"; ws.page_setup.paperSize=ws.PAPERSIZE_A4
ws.sheet_properties.pageSetUpPr=PageSetupProperties(fitToPage=True); ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=1
# ---------------- 1 Indicadores ----------------
ws=wb.create_sheet("1 Indicadores"); cab(ws,"Indicadores brutos por município","Valores de origem dos nove indicadores do PCM e variáveis auxiliares. Proporções em percentual.")
cols=[("Código IBGE",11),("Município",26),("Porte",7),("População total (2022)",12),("População de 0 a 4 anos (2022)",12),
("Matrículas em creche (2023)",12),("MC: cobertura de creche",11),("MP: matrícula pré-escola e fundamental por criança de 5 a 14 anos",14),("Estabelecimentos privados com creche",12),("Estabelecimentos com creche",12),("OP: oferta privada",11),
("MI: mortalidade infantil (por mil)",12),("PC: proporção de crianças",11),("TN: taxa de natalidade (por mil)",12),
("BP: baixo peso para a idade",11),("Crianças no CadÚnico a acompanhar",12),("CU: crianças no CadÚnico",11),("IR: IPDM renda, emprego e produção",12)]
grupos(ws,R0-2,[(1,5,"Município","BASE"),(6,11,"Fator educacional","FE"),(12,14,"Fator demográfico","FD"),(15,18,"Fator socioeconômico","FS")])
cabecalho(ws,R0-1,cols,"BASE")
base_cols(ws)
for i,row in b.iterrows():
    r=R0+i
    vals=[row["pop"],row.pop04,row.mat_creche,row.prop_creche,row.prop_pre_fund,row.oferta_priv,row.oferta_tot,row.prop_priv,row.mort_inf,row.prop04,row.natalidade,row.desnut_prop,row.cadunico,row.cad_prop,row.ipdm_r]
    for j,v in enumerate(vals,4): ws.cell(row=r,column=j,value=None if pd.isna(v) else float(v))
corpo(ws,1,18,{1:"0",4:"#,##0",5:"#,##0",6:"#,##0",7:"0.0%",8:"0.00",9:"#,##0",10:"#,##0",11:"0.0%",12:"0.0",13:"0.0%",14:"0.00",15:"0.0%",16:"#,##0",17:"0.0%",18:"0.000"})
rodape(ws,["Fontes: Inep (Censo Escolar 2023); IBGE (Censo Demográfico 2022); Ministério da Saúde (SIM, Sinasc, Sisvan e condicionalidades do Bolsa Família); Ipardes (natalidade 2022; IPDM 2021).","PC: coluna da planilha de cálculo; corresponde, pela magnitude, à faixa de 0 a 9 anos (ver aba Sobre)."])
fechar(ws,"D6")
# ---------------- 2-4 Fatores ----------------
def aba_fator(nome,titulo,k,brutos,norms,pesos,colfator,lab):
    ws=wb.create_sheet(nome); cab(ws,titulo,"Indicadores brutos, indicadores normalizados de 0 a 100 e cálculo do fator. As colunas em cor são fórmulas.")
    nb=len(brutos); c_b0=4; c_n0=c_b0+nb; c_media=c_n0+len(norms); c_f=c_media+1
    cols=[("Código IBGE",11),("Município",26),("Porte",7)]+[(t,13) for t,_,_ in brutos]+[(t,13) for t,_,_ in norms]+[(f"Média ponderada ({lab})",13),(f"{lab} (0 a 100)",12)]
    grupos(ws,R0-2,[(1,3,"Município","BASE"),(c_b0,c_n0-1,"Indicador bruto",k),(c_n0,c_media-1,"Indicador normalizado (0 a 100)",k),(c_media,c_f,"Fator",k)])
    cabecalho(ws,R0-1,cols,k); base_cols(ws)
    for i,row in b.iterrows():
        r=R0+i
        for j,(_,col,_) in enumerate(brutos): ws.cell(row=r,column=c_b0+j,value=float(row[col]))
        for j,(_,col,_) in enumerate(norms): ws.cell(row=r,column=c_n0+j,value=float(row[col]))
        termos=[f"{p}*{L(c_n0+j)}{r}" if p!=1 else f"{L(c_n0+j)}{r}" for j,p in enumerate(pesos)]
        ws.cell(row=r,column=c_media,value=f"=({'+'.join(termos)})/{sum(pesos)}")
        m=L(c_media); ws.cell(row=r,column=c_f,value=f"=({m}{r}-MIN({m}${R0}:{m}${R1}))/(MAX({m}${R0}:{m}${R1})-MIN({m}${R0}:{m}${R1}))*100")
    fm={j:fmt for j,(_,_,fmt) in enumerate(brutos,c_b0)}; fm.update({j:"0.0" for j in range(c_n0,c_f+1)})
    corpo(ws,1,c_f,fm,destaque=(c_f,))
    for r in range(R0,R1+1):
        for j in (c_media,c_f): ws.cell(row=r,column=j).fill=PatternFill("solid",fgColor=COR[k][1])
    return ws,c_f
ws2,fe_col=aba_fator("2 Fator educacional","Fator educacional (F_E)","FE",
 [("MC: cobertura de creche","prop_creche","0.0%"),("MP: matrícula pré-escola e fundamental por criança","prop_pre_fund","0.00"),("OP: oferta privada","prop_priv","0.0%")],
 [("N(MC), sentido inverso, peso 2","creche_norm",None),("N(MP), sentido inverso, peso 1","mat_norm",None),("N(OP), sentido inverso, peso 1","priv_norm",None)],[2,1,1],None,"F_E")
rodape(ws2,["F_E = normalização mínimo-máximo de [(2 × N(MC) + N(MP) + N(OP)) / 4]. Indicadores normalizados conforme as planilhas de cálculo do estudo.","Fontes: Inep (Censo Escolar 2023); IBGE (Censo Demográfico 2022)."]); fechar(ws2,"D6")
ws3,fd_col=aba_fator("3 Fator demográfico","Fator demográfico (F_D)","FD",
 [("MI: mortalidade infantil (por mil)","mort_inf","0.0"),("PC: proporção de crianças","prop04","0.0%"),("TN: taxa de natalidade (por mil)","natalidade","0.00")],
 [("N(MI), sentido direto","mi_norm",None),("N(PC), sentido direto","prop04_norm",None),("N(TN), sentido direto","nat_norm",None)],[1,1,1],None,"F_D")
rodape(ws3,["F_D = normalização mínimo-máximo de [(N(MI) + N(PC) + N(TN)) / 3].","Fontes: Ministério da Saúde (SIM e Sinasc, 2020); IBGE (Censo Demográfico 2022); Ipardes (natalidade 2022)."]); fechar(ws3,"D6")
ws4,fs_col=aba_fator("4 Fator socioeconômico","Fator socioeconômico (F_S)","FS",
 [("BP: baixo peso para a idade","desnut_prop","0.0%"),("CU: crianças no CadÚnico","cad_prop","0.0%"),("IR: IPDM renda, emprego e produção","ipdm_r","0.000")],
 [("N(BP), sentido direto","desnut_norm",None),("N(CU), sentido direto","cad_norm",None),("N(IR), sentido inverso","ipdm_norm",None)],[1,1,1],None,"F_S")
rodape(ws4,["F_S = normalização mínimo-máximo de [(N(BP) + N(CU) + N(IR)) / 3].","Fontes: Ministério da Saúde (Sisvan e condicionalidades do Bolsa Família, 2023); Ipardes (IPDM 2021)."]); fechar(ws4,"D6")
# ---------------- 5 PCM ----------------
ws=wb.create_sheet("5 PCM"); cab(ws,"Potencial de Creche por Município (PCM)","PCM = (2 × F_E + F_D + 2 × F_S) / 500. Fatores trazidos das abas 2 a 4.")
cols=[("Código IBGE",11),("Município",26),("Porte",7),("F_E",10),("F_D",10),("F_S",10),("PCM (0 a 1)",11),("Posição no estado",11),("Posição na faixa de porte",11),("Municípios na faixa",11)]
grupos(ws,R0-2,[(1,3,"Município","BASE"),(4,4,"Peso 2","FE"),(5,5,"Peso 1","FD"),(6,6,"Peso 2","FS"),(7,10,"Indicador agregado","PCM")])
cabecalho(ws,R0-1,cols,"PCM"); base_cols(ws)
for i in range(N):
    r=R0+i
    ws.cell(row=r,column=4,value=f"='2 Fator educacional'!{L(fe_col)}{r}")
    ws.cell(row=r,column=5,value=f"='3 Fator demográfico'!{L(fd_col)}{r}")
    ws.cell(row=r,column=6,value=f"='4 Fator socioeconômico'!{L(fs_col)}{r}")
    ws.cell(row=r,column=7,value=f"=(2*D{r}+E{r}+2*F{r})/500")
    ws.cell(row=r,column=8,value=f"=RANK(G{r},G${R0}:G${R1},0)")
    ws.cell(row=r,column=9,value=f"=COUNTIFS(C${R0}:C${R1},C{r},G${R0}:G${R1},\">\"&G{r})+1")
    ws.cell(row=r,column=10,value=f"=COUNTIF(C${R0}:C${R1},C{r})")
corpo(ws,1,10,{1:"0",4:"0.0",5:"0.0",6:"0.0",7:"0.000",8:"0",9:"0",10:"0"},destaque=(7,))
for r in range(R0,R1+1): ws.cell(row=r,column=7).fill=PatternFill("solid",fgColor=COR["PCM"][1])
for col,k in [(4,"FE"),(5,"FD"),(6,"FS")]:
    for r in range(R0,R1+1): ws.cell(row=r,column=col).font=F(10,color=COR[k][0])
rodape(ws,["Posição 1 = maior PCM. Posição na faixa: ordem do município entre os de mesmo porte, usada no ranqueamento das faixas P3, P2 e P1."]); fechar(ws,"D6")
# ---------------- 6 Distribuição ----------------
ws=wb.create_sheet("6 Distribuição"); cab(ws,"Creches atribuídas por município","Etapa prévia (Deliberação CEDCA 60/2023), versões da distribuição de março e maio de 2024 e lista publicada na Resolução SEDEF nº 219/2024.")
cols=[("Código IBGE",11),("Município",26),("Porte",7),("PCM",9),("Etapa prévia (índice Ipardes)",12),("Etapa PCM, março de 2024",12),("Total, março de 2024",11),("Total, maio de 2024 (versão apresentada ao CEDCA)",14),("Lista publicada, Res. SEDEF 219/2024",13),("Contemplado na lista publicada",12)]
grupos(ws,R0-2,[(1,4,"Município","BASE"),(5,7,"Março de 2024","PCM"),(8,8,"Maio de 2024","PCM"),(9,10,"Junho de 2024","PCM")])
cabecalho(ws,R0-1,cols,"BASE"); base_cols(ws)
for i,row in b.iterrows():
    r=R0+i
    ws.cell(row=r,column=4,value=f"='5 PCM'!G{r}")
    ws.cell(row=r,column=5,value=int(row.cons_prev43)); ws.cell(row=r,column=6,value=int(row.cons_pcm257))
    ws.cell(row=r,column=7,value=f"=E{r}+F{r}"); ws.cell(row=r,column=8,value=int(row.cedca_258)); ws.cell(row=r,column=9,value=int(row.pub_res219))
    ws.cell(row=r,column=10,value=f'=IF(I{r}>0,"Sim","Não")')
corpo(ws,1,10,{1:"0",4:"0.000",5:"0",6:"0",7:"0",8:"0",9:"0"})
for r in range(R0,R1+1): ws.cell(row=r,column=10).alignment=Alignment(horizontal="center")
rt=R1+1
ws.cell(row=rt,column=2,value="Paraná").font=F(10,True)
for col in range(5,10):
    c=ws.cell(row=rt,column=col,value=f"=SUM({L(col)}{R0}:{L(col)}{R1})"); c.font=F(10,True); c.number_format="0"; c.border=Border(top=forte)
c=ws.cell(row=rt,column=10,value=f'=COUNTIF(J{R0}:J{R1},"Sim")&" municípios"'); c.font=F(10,True); c.border=Border(top=forte)
for col in range(1,5): ws.cell(row=rt,column=col).border=Border(top=forte)
ws.cell(row=rt+2,column=1,value="Fontes: planilhas de cálculo do estudo (versões de 27 de março e 21 de maio de 2024); Paraná, Resolução SEDEF nº 219/2024, Anexo I.").font=F(9,color=CINZA,italic=True)
fechar(ws,"D6")
# ---------------- 7 Resumo por porte ----------------
ws=wb.create_sheet("7 Resumo por porte"); cab(ws,"Resumo por faixa de porte","Totais calculados por fórmula a partir das abas 1, 5 e 6.")
cols=[("Faixa de porte (habitantes)",26),("Municípios",11),("População de 0 a 4 anos",13),("Cobertura de creche",11),("F_E médio",10),("F_D médio",10),("F_S médio",10),("PCM médio",10),("Etapa prévia",10),("Total, março de 2024",11),("Total, maio de 2024",11),("Lista publicada",11),("Municípios contemplados",12)]
grupos(ws,4,[(1,4,"Faixa","BASE"),(5,8,"Indicador (média da faixa)","PCM"),(9,13,"Creches","BASE")])
cabecalho(ws,5,cols,"BASE")
I="'1 Indicadores'"; P="'5 PCM'"; D="'6 Distribuição'"
for k,p in enumerate(ORD):
    r=6+k
    ws.cell(row=r,column=1,value=PORTE[p])
    cond=f'"{p}"'
    ws.cell(row=r,column=2,value=f"=COUNTIF({I}!$C${R0}:$C${R1},{cond})")
    ws.cell(row=r,column=3,value=f"=SUMIF({I}!$C${R0}:$C${R1},{cond},{I}!$E${R0}:$E${R1})")
    ws.cell(row=r,column=4,value=f"=SUMIF({I}!$C${R0}:$C${R1},{cond},{I}!$F${R0}:$F${R1})/C{r}")
    for col,src in [(5,"D"),(6,"E"),(7,"F"),(8,"G")]:
        ws.cell(row=r,column=col,value=f"=AVERAGEIF({P}!$C${R0}:$C${R1},{cond},{P}!${src}${R0}:${src}${R1})")
    for col,src in [(9,"E"),(10,"G"),(11,"H"),(12,"I")]:
        ws.cell(row=r,column=col,value=f"=SUMIF({D}!$C${R0}:$C${R1},{cond},{D}!${src}${R0}:${src}${R1})")
    ws.cell(row=r,column=13,value=f'=COUNTIFS({D}!$C${R0}:$C${R1},{cond},{D}!$I${R0}:$I${R1},">0")')
rt=6+len(ORD)
ws.cell(row=rt,column=1,value="Paraná")
for col in (2,3,9,10,11,12,13): ws.cell(row=rt,column=col,value=f"=SUM({L(col)}6:{L(col)}{rt-1})")
ws.cell(row=rt,column=4,value=f"=SUM({I}!$F${R0}:$F${R1})/C{rt}")
for col,src in [(5,"D"),(6,"E"),(7,"F"),(8,"G")]: ws.cell(row=rt,column=col,value=f"=AVERAGE({P}!${src}${R0}:${src}${R1})")
fmts={2:"0",3:"#,##0",4:"0.0%",5:"0.0",6:"0.0",7:"0.0",8:"0.000",9:"0",10:"0",11:"0",12:"0",13:"0"}
for r in range(6,rt+1):
    for j in range(1,14):
        c=ws.cell(row=r,column=j); c.font=F(10,bold=(r==rt)); c.border=Border(bottom=fino) if r<rt else Border(top=forte,bottom=forte)
        if j in fmts: c.number_format=fmts[j]; c.alignment=Alignment(horizontal="right")
for r in range(6,rt+1): ws.cell(row=r,column=8).fill=PatternFill("solid",fgColor=COR["PCM"][1])
ws.cell(row=rt+2,column=1,value="Cobertura: matrículas em creche (2023) por criança de 0 a 4 anos (2022). Médias simples dos municípios da faixa.").font=F(9,color=CINZA,italic=True)
ws.freeze_panes="B6"; ws.page_setup.orientation="landscape"; ws.page_setup.paperSize=ws.PAPERSIZE_A4
ws.sheet_properties.pageSetUpPr=PageSetupProperties(fitToPage=True); ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=1
# propriedades
for w in wb.worksheets: w.sheet_view.zoomScale=100
wb.worksheets[0].sheet_properties.tabColor="3F3F3F"
for nm,k in [("2 Fator educacional","FE"),("3 Fator demográfico","FD"),("4 Fator socioeconômico","FS"),("5 PCM","PCM")]: wb[nm].sheet_properties.tabColor=COR[k][0]
wb.properties.creator=""; wb.properties.lastModifiedBy=""; wb.properties.title=""
from openpyxl.workbook.properties import CalcProperties
wb.calculation=CalcProperties(fullCalcOnLoad=True)
wb.save(OUT); print("salvo",OUT)
