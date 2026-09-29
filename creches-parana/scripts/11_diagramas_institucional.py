"""Diagramas editáveis: (a) estrutura institucional e funcionamento do Programa Infância Feliz Paraná;
(b) do PCM às creches: cálculo ponderado e arredondamento (planilha de parâmetros).
Saídas: SVG com texto editável, PDF, PNG (matplotlib) e PPTX com formas nativas (python-pptx)."""
import sys; sys.path.insert(0,"scripts")
from estilo import *
import matplotlib.pyplot as plt, textwrap
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
from pptx import Presentation
from pptx.util import Cm as PCm, Pt as PPt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
OUT="06_diagramas/"
LH=0.0353*1.25
def canvas(wcm,hcm):
    fig,ax=plt.subplots(figsize=(wcm*CM,hcm*CM)); ax.set_xlim(0,wcm); ax.set_ylim(0,hcm); ax.set_axis_off(); return fig,ax
def quebrar(texto,w,fs,bold,pad):
    cpp=fs*(0.0245 if bold else 0.0228); wrap=max(10,int((w-2*pad)/cpp)); return "\n".join(textwrap.fill(l,wrap) for l in texto.split("\n"))
def altura(texto,w,fs=7,bold=False,pad=0.25):
    t=quebrar(texto,w,fs,bold,pad); n=t.count("\n")+1; return n*fs*LH+2*pad*0.8
def box(ax,x,ytop,w,texto,fc="white",ec=TINTA,fs=7,bold=False,tc=TINTA,lw=0.6,pad=0.25,hmin=0.0,h=None):
    t=quebrar(texto,w,fs,bold,pad); n=t.count("\n")+1
    hh=h if h else max(hmin,n*fs*LH+2*pad*0.8)
    ax.add_patch(FancyBboxPatch((x,ytop-hh),w,hh,boxstyle="round,pad=0,rounding_size=0.1",fc=fc,ec=ec,lw=lw))
    ax.text(x+w/2,ytop-hh/2,t.replace("$","\\$"),ha="center",va="center",fontsize=fs,fontweight="bold" if bold else "normal",color=tc,linespacing=1.15)
    return ytop-hh
def seta(ax,x0,y0,x1,y1,color=TINTA,lw=0.6,style="-|>"):
    ax.add_patch(FancyArrowPatch((x0,y0),(x1,y1),arrowstyle=style,mutation_scale=7,color=color,lw=lw,shrinkA=0,shrinkB=0))
# ---------- registro das formas para exportar ao PPTX ----------
class Reg:
    def __init__(s): s.boxes=[]; s.arrows=[]; s.texts=[]
R=Reg()
def rbox(ax,x,ytop,w,texto,**k):
    yb=box(ax,x,ytop,w,texto,**k); R.boxes.append((x,ytop,w,ytop-yb,texto,k.get("fc","white"),k.get("ec",TINTA),k.get("tc",TINTA),k.get("fs",7),k.get("bold",False))); return yb
def rseta(ax,x0,y0,x1,y1,**k):
    seta(ax,x0,y0,x1,y1,**k); R.arrows.append((x0,y0,x1,y1,k.get("color",TINTA),k.get("style","-|>")))
def rtext(ax,x,y,t,fs=7,color=TINTA,ha="center",bold=False):
    ax.text(x,y,t,ha=ha,va="center",fontsize=fs,color=color,fontweight="bold" if bold else "normal"); R.texts.append((x,y,t,fs,color,ha,bold))
def hexrgb(h):
    h={"white":"#FFFFFF","black":"#000000"}.get(h,h).lstrip("#"); return RGBColor(int(h[0:2],16),int(h[2:4],16),int(h[4:6],16))
def exportar_pptx(reg,W,H,titulo,path):
    prs=Presentation(); prs.slide_width=PCm(W+1); prs.slide_height=PCm(H+1.6)
    sl=prs.slides.add_slide(prs.slide_layouts[6])
    tb=sl.shapes.add_textbox(PCm(0.5),PCm(0.2),PCm(W),PCm(0.8)); p=tb.text_frame.paragraphs[0]; p.text=titulo; p.font.size=PPt(11); p.font.name="Arial"; p.font.color.rgb=hexrgb(TINTA)
    off=1.1
    def Y(y): return PCm(H-y+off)
    for (x,ytop,w,h,texto,fc,ec,tc,fs,bold) in reg.boxes:
        sh=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,PCm(x+0.5),Y(ytop),PCm(w),PCm(h))
        sh.adjustments[0]=0.08; sh.fill.solid(); sh.fill.fore_color.rgb=hexrgb(fc); sh.line.color.rgb=hexrgb(ec); sh.line.width=PPt(0.75); sh.shadow.inherit=False
        tf=sh.text_frame; tf.word_wrap=True; tf.vertical_anchor=MSO_ANCHOR.MIDDLE
        for m in ("margin_left","margin_right","margin_top","margin_bottom"): setattr(tf,m,PCm(0.1))
        for i,l in enumerate(texto.split("\n")):
            pp=tf.paragraphs[0] if i==0 else tf.add_paragraph(); pp.text=l; pp.alignment=PP_ALIGN.CENTER; pp.font.size=PPt(fs); pp.font.name="Arial"; pp.font.bold=bold; pp.font.color.rgb=hexrgb(tc)
    for (x0,y0,x1,y1,color,style) in reg.arrows:
        c=sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,PCm(x0+0.5),Y(y0),PCm(x1+0.5),Y(y1)); c.line.color.rgb=hexrgb(color); c.line.width=PPt(0.75)
        ln=c.line._get_or_add_ln(); 
        from pptx.oxml.ns import qn
        from lxml import etree
        tail=etree.SubElement(ln,qn("a:tailEnd")); tail.set("type","triangle"); tail.set("w","med"); tail.set("len","med")
    for (x,y,t,fs,color,ha,bold) in reg.texts:
        w=len(t)*fs*0.022+0.4; tb=sl.shapes.add_textbox(PCm(x+0.5-(w/2 if ha=="center" else (w if ha=="right" else 0))),Y(y+0.25),PCm(w),PCm(0.5))
        pp=tb.text_frame.paragraphs[0]; pp.text=t; pp.font.size=PPt(fs); pp.font.name="Arial"; pp.font.bold=bold; pp.font.color.rgb=hexrgb(color); pp.alignment={"center":PP_ALIGN.CENTER,"left":PP_ALIGN.LEFT,"right":PP_ALIGN.RIGHT}[ha]
    prs.save(path)
# =============== Figura 1: estrutura institucional, financiamento e funcionamento ===============
R=Reg(); W,H=16,11.0; fig,ax=canvas(W,H)
ax.text(0,H-0.05,"Figura 1. Arranjo institucional e financeiro do Programa Infância Feliz Paraná",fontsize=9,color=TINTA,va="top")
cw=4.9; g=0.65
y0=H-0.7
rtext(ax,cw/2,y0-0.2,"Critério",fs=7.5,bold=True); rtext(ax,cw+g+cw/2,y0-0.2,"Deliberação e regulação",fs=7.5,bold=True); rtext(ax,2*(cw+g)+cw/2,y0-0.2,"Financiamento e execução",fs=7.5,bold=True)
y=y0-0.5
T1=["Casa Civil\nEstudo técnico: PCM, portes e regra de alocação","CEDCA/PR (conselho paritário)\nAporte do FIA e aprovação do estudo (Deliberações 60/2023 e 25/2024)","Fontes: Tesouro do Estado, FIA/PR e Assembleia Legislativa (R$ 100 milhões em 2024)"]
T2=["Ipardes\nEstudo prévio (43 unidades) e bases de dados","SEDEF (gestora do FIA/PR)\nResoluções 212 e 219/2024: critério, lista, habilitação e repasse","Paranacidade / SECID\nProjeto padrão, análise, autorização para licitar e vistorias"]
H1=max(altura(t,cw,7) for t in T1); H2=max(altura(t,cw,7) for t in T2)
rbox(ax,0,y,cw,T1[0],fc=CINZAS[1],ec=CINZAS[1],tc="white",fs=7,h=H1)
rbox(ax,cw+g,y,cw,T1[1],fc=FS_ESC,ec=FS_ESC,tc="white",fs=7,h=H1)
rbox(ax,2*(cw+g),y,cw,T1[2],fc="#F2F2F2",ec=BORDA,fs=7,h=H1)
y2=y-H1-0.45
rbox(ax,0,y2,cw,T2[0],fc="#F2F2F2",ec=BORDA,fs=7,h=H2)
rbox(ax,cw+g,y2,cw,T2[1],fc=FD_ESC,ec=FD_ESC,tc="white",fs=7,h=H2)
rbox(ax,2*(cw+g),y2,cw,T2[2],fc=FE_ESC,ec=FE_ESC,tc="white",fs=7,h=H2)
rseta(ax,cw,y-H1/2,cw+g,y-H1/2)
rseta(ax,cw+g+cw/2,y-H1,cw+g+cw/2,y2)
rseta(ax,2*(cw+g),y-H1/2,2*cw+g,y-H1/2)
rseta(ax,cw/2,y2,cw/2,y-H1)
rseta(ax,2*(cw+g)+cw/2,y-H1,2*(cw+g)+cw/2,y2)
ym=y2-H2-0.45
TM="Incentivo financeiro fundo a fundo (FIA/PR para os fundos municipais), por Termo de Adesão, em cinco parcelas condicionadas à execução; sem convênio e sem operação de crédito"
HM=altura(TM,W,7)
rbox(ax,0,ym,W,TM,fc=CINZA_CLARO,ec=CINZA_CLARO,fs=7,h=HM)
rseta(ax,cw+g+cw/2,y2-H2,cw+g+cw/2,ym)
rseta(ax,2*(cw+g)+cw/2,y2-H2,2*(cw+g)+cw/2,ym)
yk=ym-HM-0.35
TK="Municípios (399): adesão, terreno, projeto de implantação, licitação, obra e prestação de contas; custo excedente e custeio posterior por conta do município"
HK=altura(TK,W,7)
rbox(ax,0,yk,W,TK,fc="white",ec=TINTA,fs=7,h=HK)
rseta(ax,W/2,ym-HM,W/2,yk)
yc=yk-HK-0.35
rbox(ax,0,yc,7.6,"Controle: TCE-PR, MP-PR, CEDCA e conselhos municipais",fc="white",ec=BORDA,fs=7,h=0.8)
rbox(ax,8.4,yc,7.6,"Resultado: 303 unidades em 261 municípios (2024); segunda etapa em 2025",fc="white",ec=BORDA,fs=7,h=0.8)
ax.set_ylim(yc-0.8-0.2,H)
fig.text(0,0,"Fonte: Elaboração própria a partir da Lei Estadual 21.870/2023, das Deliberações CEDCA 60/2023 e 25/2024 e das Resoluções SEDEF 212 e 219/2024 e 029/2025.",fontsize=7,color=TINTA2,va="top")
salvar(fig,OUT+"figura_institucional"); plt.close(fig)
exportar_pptx(R,W,H,"Figura 1. Arranjo institucional e financeiro do Programa Infância Feliz Paraná",OUT+"figura_institucional.pptx")
# =============== Figura 3: aplicação do PCM (cotas, arredondamento e limites) ===============
R=Reg(); W,H=16,9.2; fig,ax=canvas(W,H)
ax.text(0,H-0.05,"Figura 3. Aplicação do PCM: da cota populacional às unidades inteiras",fontsize=9,color=TINTA,va="top")
y=H-0.65
y=rbox(ax,0,y,W,"Parâmetros: 257 unidades a distribuir, oito portes, população-alvo por porte, PCM por município, tetos e cotas por faixa",fc="#F2F2F2",ec=BORDA,fs=7.2)
rseta(ax,W/2,y,W/2,y-0.3); y-=0.3
wl=7.7; xr=W-wl
hh=max(altura("1. Cota do porte\nproporcional à população-alvo, corrigida pelo PCM médio (eq. 6)",wl,7),altura("2. Cota do município\nparcela do município na população-alvo do porte (eq. 7)",wl,7))
yl=rbox(ax,0,y,wl,"1. Cota do porte\nproporcional à população-alvo, corrigida pelo PCM médio (eq. 6)",fc="white",ec=CINZAS[1],fs=7,h=hh)
yr=rbox(ax,xr,y,wl,"2. Cota do município\nparcela do município na população-alvo do porte (eq. 7)",fc="white",ec=CINZAS[1],fs=7,h=hh)
rseta(ax,wl,y-hh/2,xr,y-hh/2)
y=yl-0.35; rseta(ax,xr+wl/2,yr,xr+wl/2,y)
y=rbox(ax,0,y,W,"3. Ajuste pelo indicador: K = cota × (1 + PCM) − unidades da etapa prévia (eq. 8)",fc=PCM_ESC,ec=PCM_ESC,tc="white",fs=7.2,bold=True)
rseta(ax,W/2,y,W/2,y-0.3); y-=0.3
hh=max(altura("4. Parte inteira de K (eq. 9)",wl,7),altura("5. Maior resto: unidades restantes às maiores frações (eq. 10)",wl,7))
yl=rbox(ax,0,y,wl,"4. Parte inteira de K (eq. 9)",fc="white",ec=CINZAS[1],fs=7,h=hh)
yr=rbox(ax,xr,y,wl,"5. Maior resto: unidades restantes às maiores frações (eq. 10)",fc="white",ec=CINZAS[1],fs=7,h=hh)
rseta(ax,wl,y-hh/2,xr,y-hh/2)
y=yl-0.35; rseta(ax,xr+wl/2,yr,xr+wl/2,y)
hh=max(altura("6a. Portes grandes: teto por município (10, 8, 7, 4 e 2)",wl,7),altura("6b. Portes pequenos: uma unidade aos municípios de maior PCM (27, 100 e 30)",wl,7))
yl=rbox(ax,0,y,wl,"6a. Portes grandes: teto por município (10, 8, 7, 4 e 2)",fc=CINZAS[2],ec=CINZAS[2],tc="white",fs=7,h=hh)
yr=rbox(ax,xr,y,wl,"6b. Portes pequenos: uma unidade aos municípios de maior PCM (27, 100 e 30)",fc=CINZAS[2],ec=CINZAS[2],tc="white",fs=7,h=hh)
y=yl-0.35; rseta(ax,wl/2,yl,wl/2,y); rseta(ax,xr+wl/2,yr,xr+wl/2,y)
y=rbox(ax,0,y,W,"Resultado: 43 + 100 + 157 = 300 unidades; a lista publicada (303 em 261 municípios) reduziu os tetos dos portes grandes",fc="white",ec=TINTA,fs=7.2,bold=True)
ax.set_ylim(y-0.2,H)
fig.text(0,0,"Fonte: Elaboração própria a partir das planilhas de cálculo (abas Metodologia e Dados) e da Resolução SEDEF 219/2024.",fontsize=7,color=TINTA2,va="top")
salvar(fig,OUT+"figura_calculo_ponderado"); plt.close(fig)
exportar_pptx(R,W,H,"Figura 3. Aplicação do PCM: da cota populacional às unidades inteiras",OUT+"figura_calculo_ponderado.pptx")
# =============== Figuras 4 e 5: perguntas que estruturaram o PCM e sua aplicação ===============
def fig_perguntas(nome,titulo,etapas,arquivo,fonte):
    global R
    R=Reg(); W=16; wq=6.0; wd=9.4; g=0.6; H=0.9+len(etapas)*1.45; fig,ax=canvas(W,H)
    ax.text(0,H-0.05,titulo,fontsize=9,color=TINTA,va="top")
    y=H-0.75
    rtext(ax,wq/2,y,"Pergunta",fs=7.5,bold=True); rtext(ax,wq+g+wd/2,y,"Decisão",fs=7.5,bold=True)
    y-=0.3
    for i,(tag,perg,dec,cor) in enumerate(etapas):
        hh=max(altura(tag+"\n"+perg,wq,7,pad=0.2),altura(dec,wd,7,pad=0.2),0.95)
        rbox(ax,0,y,wq,tag+"\n"+perg,fc=cor,ec=cor,tc="white",fs=7,pad=0.2,h=hh)
        rbox(ax,wq+g,y,wd,dec,fc="white",ec=cor,fs=7,pad=0.2,h=hh)
        rseta(ax,wq,y-hh/2,wq+g,y-hh/2,color=cor)
        ynext=y-hh-0.25
        if i<len(etapas)-1: rseta(ax,wq/2,y-hh,wq/2,ynext,color=cor)
        y=ynext
    ax.set_ylim(y-0.05,H)
    fig.text(0,0,fonte,fontsize=7,color=TINTA2,va="top")
    salvar(fig,OUT+arquivo); plt.close(fig)
    exportar_pptx(R,W,H,titulo,OUT+arquivo+".pptx")
fig_perguntas("f4","Figura 4. Perguntas que estruturaram a construção do PCM",[
 ("1. Necessidade","O que significa um município precisar de creche?","Três pilares: oferta educacional, pressão demográfica e vulnerabilidade socioeconômica",FE_ESC),
 ("2. Dados","Há dado oficial e recente para os 399 municípios?","Nove indicadores de bases públicas (Censo Escolar, Censo 2022, SIM/Sinasc, Sisvan, CadÚnico, IPDM), unidos pelo código do IBGE",FD_ESC),
 ("3. Comparabilidade","Como pôr indicadores de unidades distintas na mesma escala?","Normalização mínimo-máximo de 0 a 100; cobertura, oferta privada e renda em sentido inverso",FD_ESC),
 ("4. Ponderação","Quanto pesa cada indicador e cada pilar?","Cobertura de creche com peso 2 no pilar educacional; PCM = (2·FE + FD + 2·FS)/5",FS_ESC),
 ("5. Verificação","O ordenamento é defensável?","Correlação entre pilares, distribuição por porte e no mapa, comparação com o INC, sensibilidade a pesos e normalização",PCM_ESC),
],"figura_perguntas_pcm","Fonte: Elaboração própria. Cores: pilar educacional (laranja), dados e escala (azul), ponderação (verde), indicador (roxo).")
fig_perguntas("f5","Figura 5. Perguntas da aplicação: do PCM às unidades",[
 ("1. Escala","Como comparar a capital com um município de 3 mil habitantes?","Oito portes populacionais; cota de cada porte proporcional à população-alvo",CINZAS[0]),
 ("2. Conversão","Como transformar um índice contínuo em unidades inteiras?","Cota do município × (1 + PCM); parte inteira; maior resto",PCM_ESC),
 ("3. Concentração","Como evitar que poucos municípios recebam quase tudo?","Tetos por município nos portes grandes; uma unidade aos municípios de maior PCM nos pequenos",CINZAS[1]),
 ("4. Validação","Quem aprova o critério e como ele vira lista?","Deliberação do CEDCA; Resolução SEDEF com critério e lista; o próximo ranqueado substitui quem não adere",CINZAS[2]),
 ("5. Limites","O que o índice não capta?","Terreno, capacidade de licitar e fila local; margem de decisão delimitada e documentada",ACENTO),
],"figura_perguntas_aplicacao","Fonte: Elaboração própria.")
print("institucional ok")
# =============== Figura: mecanismo (dados -> PCM -> parâmetros -> deliberação -> lista -> alocação) ===============
R=Reg(); W,H=16,5.6; fig,ax=canvas(W,H)
ax.text(0,H-0.05,"Figura 6. Mecanismo: das evidências municipais à alocação do investimento",fontsize=9,color=TINTA,va="top")
et=[("Dados municipais\nnove indicadores oficiais",CINZA_CLARO,TINTA),("PCM\nmín-máx, três pilares, pesos 2:1:2",PCM_ESC,"white"),("Faixas e parâmetros\nportes, cotas, tetos, maior resto",CINZAS[1],"white"),("Deliberação\nCEDCA aprova estudo e ranqueamento",FS_ESC,"white"),("Lista de municípios\nRes. SEDEF 219/2024",FD_ESC,"white"),("Alocação do investimento\nadesão, repasse, obra",CINZAS[2],"white")]
w=2.35; g=(W-6*w)/5; y=H-0.75; hh=max(altura(t,w,7,pad=0.2) for t,_,_ in et)
for k,(t,fc,tc) in enumerate(et):
    x=k*(w+g); rbox(ax,x,y,w,t,fc=fc,ec=fc if fc!=CINZA_CLARO else BORDA,tc=tc,fs=7,pad=0.2,h=hh)
    if k<5: rseta(ax,x+w,y-hh/2,x+w+g,y-hh/2)
ax.set_ylim(y-hh-0.3,H)
fig.text(0,0,"Fonte: Elaboração própria.",fontsize=7,color=TINTA2,va="top")
salvar(fig,OUT+"figura_mecanismo"); plt.close(fig)
exportar_pptx(R,W,H,"Figura 6. Mecanismo: das evidências municipais à alocação do investimento",OUT+"figura_mecanismo.pptx")
