"""Gera a nota técnica do contrafactual populacional (material suplementar), sem identificação de autoria.
Uso: python3.13 -I nota.py <numeros.json> <saida.docx>"""
import sys, json
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

N = json.load(open(sys.argv[1], encoding='utf8'))
def br(x, d=1):
    return f"{x:.{d}f}".replace('.', ',')
def mil(x):
    return f"{x:,}".replace(',', '.')
M, Y, C, C43 = N['marco'], N['maio'], N['contraf'], N['contraf43']
L = N['limiar']

doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
sec.top_margin = Cm(1.8); sec.bottom_margin = Cm(1.6)
sec.left_margin = sec.right_margin = Cm(2.5)
st = doc.styles['Normal']
st.font.name = 'Times New Roman'; st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
st.paragraph_format.space_after = Pt(4)
st.paragraph_format.line_spacing = 1.04
cp = doc.core_properties
cp.author = ''; cp.last_modified_by = ''; cp.title = ''; cp.comments = ''; cp.keywords = ''; cp.subject = ''

def par(txt, bold=False, italic=False, size=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=None, before=None):
    p = doc.add_paragraph(); p.alignment = align
    if after is not None: p.paragraph_format.space_after = Pt(after)
    if before is not None: p.paragraph_format.space_before = Pt(before)
    r = p.add_run(txt); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    return p

par('NOTA TÉCNICA', bold=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
par('Análise contrafactual: distribuição das unidades sob critério apenas populacional', bold=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
par('Material suplementar do artigo “De indicador a regra de alocação: Potencial de Creche por Município no Programa Infância Feliz”', italic=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)

par('1. Pergunta', bold=True, before=2)
par('A seção 4.2 do artigo compara a distribuição aprovada para o Programa Infância Feliz com a que se observaria na ausência da regra adotada (o PCM e sua conversão em unidades), caso as mesmas unidades fossem repartidas apenas pela população infantil, o critério mais simples de demanda. '
    'Esta nota descreve como construímos essa simulação, em que condições as duas distribuições são comparáveis e que inferências a comparação permite.')

par('2. Construção do contrafactual', bold=True, before=2)
par(f'O contrafactual reparte as 300 unidades da versão de maio de 2024 entre os 399 municípios do Paraná em proporção à população de zero a quatro anos do Censo Demográfico de 2022, que soma {mil(L["pop04_total"])} crianças (IBGE, 2023). '
    'A cota exata do município i é qᵢ = 300 · aᵢ / A, em que aᵢ é sua população de zero a quatro anos e A, a do estado. '
    'Cada município recebe a parte inteira de qᵢ, e as unidades restantes vão às maiores partes fracionárias, pelo método do maior resto (BALINSKI; YOUNG, 2001), o mesmo que a regra adotada usa nos municípios grandes. '
    'A simulação não usa o PCM, as faixas de porte, as cotas, os tetos nem a etapa prévia. '
    f'Cada unidade corresponde, em média, a {mil(round(L["criancas_por_unidade"]))} crianças, e só recebem unidade os municípios com cota exata acima de 0,38, o equivalente a cerca de {mil(L["pop04_min_contemplado"])} crianças de zero a quatro anos.')

par('3. Comparabilidade', bold=True, before=2)
par('Mantivemos fixos os elementos que não dependem da regra: o total de unidades, o universo de municípios e a fonte populacional, o Censo Demográfico de 2022. '
    'Duas diferenças entre as distribuições comparadas precisam ser consideradas. '
    'A versão de maio inclui 43 unidades da etapa prévia, que o CEDCA/PR destinou em dezembro de 2023, antes do índice (CEDCA, 2023). '
    'Para que essas unidades não se confundam com o efeito da regra, calculamos uma segunda variante, que as mantém e reparte as outras 257 pela população. '
    'Além disso, a versão adotada usou, no indicador PC do índice e nas cotas dos municípios grandes, a coluna da planilha que contém a população de 5 a 14 anos (seção 3.2 do artigo), enquanto o contrafactual usa a de zero a quatro anos, faixa indicada pela Resolução SEDEF nº 219/2024 (PARANÁ, 2024). '
    'Com a população de zero a quatro anos e as demais correções da nota 5 do artigo, o índice teria correlação de postos de 0,99 com o adotado, 153 dos 157 municípios selecionados nas faixas menores se manteriam e duas unidades mudariam de município entre os 30 maiores, diferenças pequenas diante das que a Tabela 1 registra.')

par('4. Resultados', bold=True, before=2)
cap = par('Tabela 1 – Distribuição das 300 unidades: versões do programa e contrafactual populacional', size=10, align=WD_ALIGN_PARAGRAPH.LEFT, after=2)
cap.paragraph_format.keep_with_next = True
rows = [
    ('Indicador', 'Março de 2024', 'Maio de 2024 (aprovada)', 'Contrafactual', 'Contrafactual com etapa prévia'),
    ('Municípios contemplados', M['municipios'], Y['municipios'], C['municipios'], C43['municipios']),
    ('Municípios de até 20 mil habitantes contemplados', M['ate20mil_mun'], Y['ate20mil_mun'], C['ate20mil_mun'], C43['ate20mil_mun']),
    ('Municípios de até 5 mil habitantes contemplados', M['p1_mun'], Y['p1_mun'], C['p1_mun'], C43['p1_mun']),
    ('Unidades em Curitiba', M['curitiba'], Y['curitiba'], C['curitiba'], C43['curitiba']),
    ('Unidades nos 10 maiores receptores (%)', br(M['top10_pct']), br(Y['top10_pct']), br(C['top10_pct']), br(C43['top10_pct'])),
    ('Unidades nos 80 municípios de maior PCM (%)', br(M['top80pcm_pct']), br(Y['top80pcm_pct']), br(C['top80pcm_pct']), br(C43['top80pcm_pct'])),
    ('PCM médio ponderado pelas unidades', br(M['pcm_medio_unid'], 3), br(Y['pcm_medio_unid'], 3), br(C['pcm_medio_unid'], 3), br(C43['pcm_medio_unid'], 3)),
    ('Crianças de 0 a 4 anos em municípios contemplados (%)', br(M['cob_pop04_pct']), br(Y['cob_pop04_pct']), br(C['cob_pop04_pct']), br(C43['cob_pop04_pct'])),
    ('Municípios também contemplados na versão de maio', M['comuns_maio'], Y['comuns_maio'], C['comuns_maio'], C43['comuns_maio']),
]
t = doc.add_table(rows=len(rows), cols=5); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = 'Table Grid'
t.autofit = False
W = [Cm(6.4), Cm(2.2), Cm(2.6), Cm(2.3), Cm(2.5)]
for i, r in enumerate(rows):
    for j, v in enumerate(r):
        c = t.cell(i, j); c.width = W[j]
        p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(str(v)); run.font.size = Pt(9); run.bold = (i == 0)
grid = t._tbl.tblGrid
for j, gc in enumerate(grid.findall(qn('w:gridCol'))):
    gc.set(qn('w:w'), str(int(W[j].twips)))
for j, col in enumerate(t.columns):
    col.width = W[j]
for i, row in enumerate(t.rows):
    trPr = row._tr.get_or_add_trPr()
    cs = OxmlElement('w:cantSplit'); trPr.append(cs)
    if i < len(t.rows) - 1:
        for c in row.cells:
            for pp in c.paragraphs:
                pp.paragraph_format.keep_with_next = True
par(f'Fonte: elaboração própria com dados do Censo Demográfico de 2022 (IBGE, 2023), das planilhas do programa e da Resolução SEDEF nº 219/2024. A população de 0 a 4 anos do estado é de {mil(L["pop04_total"])} crianças.', size=9, align=WD_ALIGN_PARAGRAPH.LEFT, after=6)

par(f'Com as mesmas 300 unidades, a versão aprovada contemplou {Y["municipios"]} municípios, e o contrafactual, {C["municipios"]}. '
    f'A repartição populacional concentraria as unidades: os dez maiores receptores ficariam com {br(C["top10_pct"])}% delas, contra {br(Y["top10_pct"])}% em maio, e Curitiba receberia {C["curitiba"]} unidades, contra {Y["curitiba"]}. '
    f'Nenhum município de até 5 mil habitantes seria contemplado, contra {Y["p1_mun"]} na versão aprovada. '
    f'A regra também dirige as unidades para onde o índice mede maior necessidade: os 80 municípios de maior PCM ficaram com {br(Y["top80pcm_pct"])}% das unidades, contra {br(C["top80pcm_pct"])}% no contrafactual. '
    f'A cobertura populacional não opõe as duas distribuições. No contrafactual, {br(C["cob_pop04_pct"])}% das crianças de zero a quatro anos viveriam em municípios contemplados, contra {br(Y["cob_pop04_pct"])}% em maio e {br(M["cob_pop04_pct"])}% em março. '
    f'Com as 43 unidades da etapa prévia preservadas, as diferenças persistem ({C43["municipios"]} municípios, {C43["curitiba"]} unidades em Curitiba e {br(C43["top80pcm_pct"])}% das unidades nos 80 de maior PCM), o que afasta a etapa prévia como explicação. '
    'As diferenças restantes decorrem do conjunto da versão de maio (índice, faixas, cotas, tetos e ajustes posteriores), e a simulação não separa o peso de cada componente.')

par('5. Alcance e limites da inferência', bold=True, before=2)
par('A comparação mede o efeito alocativo da versão adotada: quantos municípios, de que porte e com que necessidade medida recebem as unidades, para um mesmo total. '
    'É nesse resultado que se apoiam as duas primeiras lições da seção 5 do artigo, segundo as quais índice e regra decidem juntos o resultado territorial e a informação intersetorial do índice muda o destino das unidades. '
    'A simulação é estática e não incorpora adesões, desistências, disponibilidade de terreno nem capacidade de execução, que na política real filtram a lista de contemplados. '
    'Também não estima efeitos sobre matrículas, emprego dos cuidadores ou desenvolvimento infantil, que dependem do funcionamento das unidades e cabem a uma avaliação ex post. '
    'O contrafactual considera só o número de crianças, sem a oferta instalada nem a vulnerabilidade, e serve por isso de referência mínima. '
    'Os dados municipais usados (população, PCM e unidades por versão) estão na planilha de dados do material suplementar, o que permite reproduzir a simulação e testar outras regras.')

par('Referências', bold=True, before=4)
for ref in [
    'BALINSKI, M. L.; YOUNG, H. P. Fair representation: meeting the ideal of one man, one vote. 2. ed. Washington, DC: Brookings Institution Press, 2001.',
    'CEDCA. Conselho Estadual dos Direitos da Criança e do Adolescente do Paraná. Deliberação nº 60/2023, de 4 de dezembro de 2023. Aprova o repasse de recursos fundo a fundo para a construção de creches. Diário Oficial do Estado do Paraná, Curitiba, 6 dez. 2023.',
    'IBGE. Censo Demográfico 2022: população por idade e sexo, resultados do universo (tabela 9514). Rio de Janeiro: IBGE, 2023. Disponível em: https://sidra.ibge.gov.br/tabela/9514. Acesso em: 29 set. 2026.',
    'PARANÁ. Secretaria de Estado do Desenvolvimento Social e Família. Resolução nº 219/2024, de 4 de junho de 2024 (retificada). Estabelece os critérios para ranqueamento de municípios do Programa Infância Feliz Paraná. Curitiba: SEDEF, 2024.',
]:
    p = par(ref, size=9, align=WD_ALIGN_PARAGRAPH.LEFT, after=2)
z = doc.settings.element.find(qn('w:zoom'))
if z is not None and z.get(qn('w:percent')) is None:
    z.set(qn('w:percent'), '100')
doc.save(sys.argv[2])
print('ok', sys.argv[2])
