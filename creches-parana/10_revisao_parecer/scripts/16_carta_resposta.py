"""Gera a carta de resposta ao parecer (Word) a partir de um JSON com os blocos de texto.
Uso: python3.13 -I carta.py <conteudo.json> <saida.docx>"""
import sys, json
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

c = json.load(open(sys.argv[1], encoding='utf8'))
doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
sec.top_margin = sec.bottom_margin = Cm(2.5)
sec.left_margin, sec.right_margin = Cm(3.0), Cm(2.0)

st = doc.styles['Normal']
st.font.name = 'Arial'; st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 1.15

def par(text, bold=False, align=None, size=None, space_after=None, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    if align == 'center': p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == 'justify': p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if space_after is not None: p.paragraph_format.space_after = Pt(space_after)
    return p

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), fill)
    tcPr.append(shd)

def cell_text(cell, text, bold=False, size=9.5, color=None):
    cell.text = ''
    paras = text.split('\n')
    for i, t in enumerate(paras):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(t); r.bold = bold; r.font.size = Pt(size)
        if color: r.font.color.rgb = RGBColor.from_string(color)

par(c['titulo'], bold=True, align='center', size=12, space_after=4)
par(c['subtitulo'], align='center', size=10, space_after=14, italic=True)
for t in c['abertura']:
    par(t, align='justify')

tab = doc.add_table(rows=1, cols=3)
tab.alignment = WD_TABLE_ALIGNMENT.CENTER
tab.style = 'Table Grid'
tab.autofit = False
larg = [Cm(4.2), Cm(9.2), Cm(2.6)]
grid = tab._tbl.tblGrid
for gc, wd in zip(grid.findall(qn('w:gridCol')), larg):
    gc.set(qn('w:w'), str(int(wd.twips)))
tblPr = tab._tbl.tblPr
tw = tblPr.find(qn('w:tblW'))
if tw is None:
    tw = OxmlElement('w:tblW'); tblPr.append(tw)
tw.set(qn('w:type'), 'dxa'); tw.set(qn('w:w'), str(int(sum(x.twips for x in larg))))
# tab.autofit = False já grava w:tblLayout fixo
for i, h in enumerate(['Ponto do parecer', 'Resposta e alteração realizada', 'Localização']):
    cl = tab.rows[0].cells[i]; cell_text(cl, h, bold=True, size=9.5, color='FFFFFF'); shade(cl, '0171C1'); cl.width = larg[i]
for k, linha in enumerate(c['quadro']):
    row = tab.add_row().cells
    cell_text(row[0], linha['ponto']); cell_text(row[1], linha['resposta']); cell_text(row[2], linha['local'])
    for i in range(3):
        row[i].width = larg[i]
        if k % 2 == 1: shade(row[i], 'EDEDED')
doc.add_paragraph()
par(c['outras_titulo'], bold=True, space_after=4)
for t in c['outras']:
    p = doc.add_paragraph(style='List Bullet'); r = p.add_run(t); r.font.size = Pt(10.5)
    p.paragraph_format.space_after = Pt(3)
for t in c.get('fecho', []):
    par(t, align='justify')
z = doc.settings.element.find(qn('w:zoom'))
if z is not None and z.get(qn('w:percent')) is None:
    z.set(qn('w:percent'), '100')
cp = doc.core_properties
cp.author = ''; cp.last_modified_by = ''; cp.title = ''; cp.comments = ''; cp.keywords = ''; cp.subject = ''
doc.save(sys.argv[2])
print('ok', sys.argv[2])
